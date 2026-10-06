"""Reads a BCM coverage sheet into clean rows.

Which header holds which field is account configuration (`AccountConfig.dp_columns`), so a renamed
column is a settings change, not a code change. Headers match ignoring case and spacing. Missing
required columns stop the import with a message naming them; missing optional ones are warnings.
"""

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

import openpyxl

from app.core.account_config import DP_FIELD_LABELS, DP_FIELDS, AccountConfig

HEADER_SEARCH_ROWS = 10  # sheets sometimes carry a title above the header


class SheetError(ValueError):
    pass


@dataclass
class ParsedRow:
    row_number: int  # as Excel shows it
    values: dict[str, Any]  # field → cleaned value
    raw: dict[str, Any]  # header → original cell, JSON-safe


@dataclass
class ParsedSheet:
    rows: list[ParsedRow]
    warnings: list[str] = field(default_factory=list)
    sheet_name: str = ""


def _norm(h: Any) -> str:
    return re.sub(r"\s+", " ", str(h or "")).strip().casefold()


def _json_safe(v: Any) -> Any:
    if isinstance(v, datetime | date):
        return v.isoformat()
    return v


def _text(v: Any, blanks: set[str]) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"[ \t]+", " ", str(v)).strip()
    return None if s == "" or s in blanks else s


def _date(v: Any, blanks: set[str]) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _text(v, blanks)
    if s is None:
        return None
    for fmt in ("%Y-%m-%d", "%d-%b-%Y", "%d-%b-%y", "%m/%d/%Y", "%d/%m/%Y", "%d %b %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise SheetError(f"'{s}' isn't a date I can read")


DATE_FIELDS = {"start_date", "doj"}
UPPER_FIELDS = {"req_id", "grade", "region"}


def parse(data: bytes, cfg: AccountConfig) -> ParsedSheet:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises several types for non-xlsx input
        raise SheetError("This isn't an Excel .xlsx file the app can read.") from e
    try:
        return _parse_workbook(wb, cfg)
    finally:
        wb.close()


def _parse_workbook(wb: Any, cfg: AccountConfig) -> ParsedSheet:
    wanted = {f: _norm(cfg.dp_columns.get(f, default)) for f, (default, _) in DP_FIELDS.items()}
    req_header = wanted["req_id"]

    for ws in wb.worksheets:
        head = list(ws.iter_rows(min_row=1, max_row=HEADER_SEARCH_ROWS, values_only=True))
        for idx, row in enumerate(head, start=1):
            normed = [_norm(c) for c in row]
            if req_header in normed:
                return _read(ws, idx, list(row), normed, wanted, cfg)
    raise SheetError(
        f"No header row with '{cfg.dp_columns.get('req_id', DP_FIELDS['req_id'][0])}' in the first "
        f"{HEADER_SEARCH_ROWS} rows of any sheet. Check the file, or the column names in Account settings."
    )


def _read(
    ws: Any,
    header_row: int,
    headers: list[Any],
    normed: list[str],
    wanted: dict[str, str],
    cfg: AccountConfig,
) -> ParsedSheet:
    col: dict[str, int] = {}
    missing_required, missing_optional = [], []
    for f, header in wanted.items():
        if header in normed:
            col[f] = normed.index(header)
        elif DP_FIELDS[f][1]:
            missing_required.append(f"{DP_FIELD_LABELS[f]} ('{cfg.dp_columns.get(f)}')")
        else:
            missing_optional.append(DP_FIELD_LABELS[f])
    if missing_required:
        raise SheetError("The sheet is missing required columns: " + ", ".join(missing_required) + ".")

    blanks = set(cfg.dp_blank_values)
    sheet = ParsedSheet(rows=[], sheet_name=ws.title)
    if missing_optional:
        sheet.warnings.append("Columns not found (left empty): " + ", ".join(missing_optional) + ".")

    for n, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        values: dict[str, Any] = {}
        for f, i in col.items():
            cell = row[i] if i < len(row) else None
            try:
                v = _date(cell, blanks) if f in DATE_FIELDS else _text(cell, blanks)
            except SheetError as e:
                sheet.warnings.append(f"Row {n}, {DP_FIELD_LABELS[f]}: {e}; left empty.")
                v = None
            if isinstance(v, str) and f in UPPER_FIELDS:
                v = v.upper().replace(" ", "")
            values[f] = v
        raw = {str(h): _json_safe(row[i]) for i, h in enumerate(headers) if h is not None and i < len(row)}
        sheet.rows.append(ParsedRow(row_number=n, values=values, raw=raw))
    if not sheet.rows:
        raise SheetError("The sheet has a header row but no data rows.")
    return sheet
