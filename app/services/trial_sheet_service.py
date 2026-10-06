"""A BCM sheet to try the import with, built from what the account holds right now.

Every requisition the app knows gets a row, most of them one step further on than the app has them, so
uploading the sheet moves demands along the workflow. One requisition is left out on purpose (it becomes
"Removed from sheet" when the previous sheet had it) and three rows belong to no demand (they are escalated
as "in the sheet, not in the
app"). The second worksheet says, row by row, what the upload should do. Candidate names are
placeholders; nothing here comes from a real sheet.
"""

import io
from dataclasses import dataclass
from datetime import date, datetime, timedelta

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.account_config import DP_FIELDS, AccountConfig
from app.core.enums import DemandStatus
from app.models import Account, Demand, OfferApproval

S = DemandStatus

# Where the sheet puts a requisition, by where the app has it today.
ADVANCE: dict[DemandStatus, tuple[DemandStatus, str]] = {
    S.SENT_TO_GTD: (
        S.COVERAGE_REQUIRED,
        "GTD has approved it: it appears in the sheet and moves to sourcing",
    ),
    S.LINKED: (S.COVERAGE_REQUIRED, "Its sheet status is now a mapped one: moves to sourcing"),
    S.MISSING: (S.COVERAGE_REQUIRED, "Back in the sheet: no longer overdue, moves to sourcing"),
    S.INTERVIEWING: (S.COVERAGE_REQUIRED, "The sheet is behind the panel: the app keeps it at the panel"),
    S.PANEL_SELECTED: (S.COVERAGE_REQUIRED, "The sheet is behind the panel: the app keeps it as selected"),
    S.PROFILES_WITH_CLIENT: (S.OFFER_IN_PROCESS, "Client selected: an offer approval is raised"),
    S.OFFER_IN_PROCESS: (S.OFFER_IN_MARKET, "Offer made: a date of joining is set"),
    S.OFFER_IN_MARKET: (S.STAFFED, "The candidate has joined: the position is fulfilled"),
    S.STAFFED: (S.STAFFED, "Still joined: nothing changes"),
    S.INCORRECT: (S.INCORRECT, "Still marked incorrect: nothing changes"),
    S.CANCELLED: (S.CANCELLED, "Still cancelled: nothing changes"),
}


# Requisitions somebody raised on GTD without a demand in the app.
OUTSIDE = [
    ("TR1AL7", "Trial: Treasury Reporting Analyst (raised outside the app)"),
    ("TR1AL8", "Trial: Collections Platform Support Engineer (raised outside the app)"),
    ("TR1AL9", "Trial: Merchant Onboarding Business Analyst (raised outside the app)"),
]


@dataclass
class TrialSheet:
    filename: str
    data: bytes
    rows: int


def _text(cfg: AccountConfig, stage: DemandStatus) -> tuple[str, str] | None:
    """How this account's sheet writes a stage: (status, status group)."""
    for m in cfg.status_mapping:
        if m.stage is stage:
            return m.status, m.status_group
    return None


def build(db: Session, account_id: int, today: date) -> TrialSheet:
    account = db.get_one(Account, account_id)
    cfg = account.settings
    head = {f: cfg.dp_columns.get(f, default) for f, (default, _) in DP_FIELDS.items()}
    blank = cfg.dp_blank_values[0] if cfg.dp_blank_values else None
    demands = [
        d
        for d in db.scalars(select(Demand).where(Demand.account_id == account_id).order_by(Demand.app_ref))
        if d.gtd_req_id and d.status_enum not in (S.CLOSED, S.DROPPED)
    ]
    sourcing = [d for d in demands if d.status_enum is S.COVERAGE_REQUIRED]
    waiting = set(db.scalars(select(OfferApproval.demand_id).where(OfferApproval.decision.is_(None))))
    leave_out = sourcing[-1] if len(sourcing) > 1 else None

    rows: list[dict[str, object]] = []
    notes: list[tuple[str, str, str]] = []
    for d in demands:
        if d is leave_out:
            notes.append(
                (
                    d.gtd_req_id or "",
                    d.name,
                    "Left out of the sheet: the previous sheet had it, so it becomes "
                    "Removed from sheet and is escalated",
                )
            )
            continue
        now = d.status_enum
        if now is S.COVERAGE_REQUIRED:
            first = d is sourcing[0]
            to, why = (
                (S.PROFILES_WITH_CLIENT, "Profiles are with the client: moves to client interview")
                if first
                else (S.COVERAGE_REQUIRED, "Still sourcing: nothing changes")
            )
        elif now is S.OFFER_IN_PROCESS and d.id in waiting:
            to, why = S.OFFER_IN_PROCESS, "Its offer still waits for approval: nothing changes"
        else:
            to, why = ADVANCE.get(now, (S.COVERAGE_REQUIRED, "In the sheet as sourcing"))
        text = _text(cfg, to) or _text(cfg, S.COVERAGE_REQUIRED)
        if text is None:
            continue
        doj: date | None = None
        if to is S.OFFER_IN_MARKET:
            doj = today + timedelta(days=14)
        elif to is S.STAFFED:
            doj = today - timedelta(days=1)
        rows.append(
            {
                "originator": d.owner.name,
                "req_id": d.gtd_req_id,
                "demand_request_name": d.name,
                "practice": d.practice,
                "grade": d.grade,
                "region": d.region,
                "start_date": datetime.combine(d.start_date, datetime.min.time()) if d.start_date else blank,
                "source": blank,
                "candidate_name": f"Candidate {d.app_ref[-3:]}"
                if to in (S.OFFER_IN_PROCESS, S.OFFER_IN_MARKET, S.STAFFED)
                else blank,
                "doj": datetime.combine(doj, datetime.min.time()) if doj else None,
                "status": text[0],
                "status_group": text[1] or blank,
            }
        )
        notes.append((d.gtd_req_id or "", d.name, f"{now.label} → {why}"))
    outside = _text(cfg, S.COVERAGE_REQUIRED)
    if outside:
        for k, (req, name) in enumerate(OUTSIDE):
            rows.append(
                {
                    "originator": "Raised outside the app",
                    "req_id": req,
                    "demand_request_name": name,
                    "practice": cfg.practices[k % len(cfg.practices)] if cfg.practices else blank,
                    "grade": cfg.grades[k % len(cfg.grades)] if cfg.grades else blank,
                    "region": blank,
                    "start_date": datetime.combine(today + timedelta(days=30 + 7 * k), datetime.min.time()),
                    "status": outside[0],
                    "status_group": outside[1] or blank,
                }
            )
            notes.append((req, name, "Not in the app: escalated to the GTD admin team"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append([f"{account.name} BCM coverage sheet (trial, generated by the app on {today:%d %b %Y})"])
    fields = list(DP_FIELDS)
    ws.append([head[f] for f in fields])
    for r in rows:
        ws.append([r.get(f, blank) for f in fields])
    guide = wb.create_sheet("what this upload does")
    guide.append(["Requisition", "Demand", "What should happen when this sheet is imported"])
    for n in notes:
        guide.append(list(n))
    for sheet, widths in ((ws, [22] * len(fields)), (guide, [14, 52, 80])):
        for i, w in enumerate(widths, start=1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return TrialSheet(
        filename=f"BCM_sheet_trial_{today.day}-{today:%b-%Y}.xlsx", data=buf.getvalue(), rows=len(rows)
    )
