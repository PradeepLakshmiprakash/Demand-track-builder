"""The trial BCM sheet: built from the account's demands, and importable as it comes."""

import io
from datetime import date

import openpyxl
from sqlalchemy.orm import Session

from app.core.enums import DemandStatus
from app.models import Demand
from app.services import trial_sheet_service
from seed import sample_sheet
from tests.conftest import Client


def test_trial_sheet_downloads_and_imports(client: Client, db: Session) -> None:
    c = client.as_user("kavya")
    assert "Download the trial sheet" in c.get("/imports").text
    # a first sheet that holds every requisition, so the one the trial sheet leaves out counts as removed
    first = c.post("/imports", files={"file": ("BCM_1-Jan-2099.xlsx", sample_sheet.build())})
    assert first.status_code == 303
    r = c.get("/imports/trial-sheet")
    assert r.status_code == 200 and "BCM_sheet_trial_" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    data, guide = wb.worksheets
    reqs = [row[0] for row in data.iter_rows(min_row=3, values_only=True)]
    assert "TR1AL9" in reqs and len(reqs) == len(set(reqs)) and guide.max_row >= len(reqs)
    # it goes through the import as it is, and moves demands on
    up = c.post("/imports", files={"file": ("BCM_sheet_trial_6-Oct-2099.xlsx", r.content)})
    assert up.status_code in (200, 303), up.text[:300]
    assert up.status_code == 303, "the trial sheet should import without an error"
    db.expire_all()
    left_out = next(row[0] for row in guide.iter_rows(min_row=2, values_only=True) if "Left out" in row[2])
    flagged = [
        d
        for d in db.query(Demand).filter(Demand.status == DemandStatus.DROPPED.value)
        if d.gtd_req_id == left_out
    ]
    assert len(flagged) == 1  # the requisition the sheet leaves out is flagged


def test_trial_sheet_leaves_one_out_and_adds_one(db: Session) -> None:
    sheet = trial_sheet_service.build(db, 1, date.today())
    wb = openpyxl.load_workbook(io.BytesIO(sheet.data))
    notes = [row[2] for row in wb.worksheets[1].iter_rows(min_row=2, values_only=True)]
    assert sum("Left out of the sheet" in n for n in notes) == 1
    assert sum("Not in the app" in n for n in notes) == 1
    assert sheet.rows > 5


def test_demand_owners_cannot_download_it(client: Client) -> None:
    assert client.as_user("priya").get("/imports/trial-sheet").status_code == 403
